"""Generate the xlsx export from the jobs DB."""

from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.db import grouped_jobs

COLUMNS = [
    ("Title", "title"),
    ("Company", "company"),
    ("Location", "location"),
    ("Status", "status"),
    ("Tier", "tier"),
    ("Years Exp", "years_experience"),
    ("Sources", "sources"),
    ("Posted", "posted_date"),
    ("First Seen", "first_seen"),
    ("Last Seen", "last_seen"),
    ("Link", "primary_url"),
]


def generate_xlsx(conn: sqlite3.Connection, output_path: str | Path) -> Path:
    """Write the grouped jobs table to an .xlsx file. Returns the path written."""
    output_path = Path(output_path)
    rows = grouped_jobs(conn)

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"

    headers = [label for label, _ in COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        primary_url = row["urls"][0][1] if row["urls"] else ""
        ws.append(
            [
                row["title"],
                row["company"],
                row["location"],
                row["status"],
                row["tier"] or "",
                row["years_experience"] if row["years_experience"] is not None else "",
                ", ".join(row["sources"]),
                row["posted_date"] or "",
                row["first_seen"],
                row["last_seen"],
                primary_url,
            ]
        )

    # Autofilter + table styling over the full data range
    last_col = get_column_letter(len(COLUMNS))
    last_row = len(rows) + 1
    table_ref = f"A1:{last_col}{last_row}"
    if len(rows) > 0:
        table = Table(displayName="Jobs", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        ws.add_table(table)

    # Reasonable column widths
    widths = [40, 28, 24, 10, 8, 10, 22, 12, 20, 20, 50]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
