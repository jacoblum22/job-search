"""Structured resume bases: load, validate and render to .docx / .pdf.

A base resume is a YAML file (kept in the git-ignored ``materials/``) whose
sections and bullets carry stable ids, so later steps can address a single
bullet (``experience/algaecal/algaecal.sql-daily``) and cite the vault
passages behind it. See ``docs/resume-format.md`` for the schema.

Template choices follow the OpenResume parse check (research note,
Oct 2026): real U+2022 bullets typed as text (not Word's Symbol-font list
bullet, which extracts as U+F0B7), "Vancouver, BC" on the contact line, the
title, company and location of each entry in separate text runs with no
pipes, ALL CAPS section headings, one column, one page.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
CONTACT_PATH = ROOT / "config" / "contact.yaml"

SECTION_TITLES = {
    "summary": "Summary",
    "skills": "Skills",
    "experience": "Experience",
    "projects": "Projects",
    "education": "Education",
}
DEFAULT_ORDER = ["summary", "skills", "experience", "projects", "education"]
ID_RE = re.compile(r"^[a-z0-9][a-z0-9.\-]*$")
BULLET = "\u2022"

FONT_NAME = "Calibri"
LINK_COLOR = "0563C1"


@dataclass
class Layout:
    """Sizes in points / inches; a base can override any of them under ``layout:``."""

    body_pt: float = 10.5
    name_pt: float = 16
    contact_pt: float = 9
    side_margin_in: float = 0.6
    top_margin_in: float = 0.5
    bottom_margin_in: float = 0.5
    section_space_pt: float = 6
    entry_space_pt: float = 3


# ---------------------------------------------------------------------------
# Loading and validation
# ---------------------------------------------------------------------------


def load_resume(path: Path) -> dict[str, Any]:
    data = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise ValueError(f"{path}: top level must be a mapping")
    return data


def load_contact(path: Path = CONTACT_PATH) -> dict[str, str]:
    """Contact details from the git-ignored ``config/contact.yaml``."""
    if not path.exists():
        raise FileNotFoundError(
            f"Missing {path}. Copy config/contact.example.yaml to config/contact.yaml "
            "and fill in your details."
        )
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    keys = ("name", "location", "email", "phone", "linkedin", "github")
    return {k: str(data.get(k) or "").strip() for k in keys}


def iter_bullets(data: dict[str, Any]) -> Iterator[tuple[str, dict[str, Any]]]:
    """Yield ``(path, bullet)`` for every bullet, e.g. ``experience/algaecal/algaecal.sql``."""
    for section in ("experience", "projects"):
        for entry in data.get(section) or []:
            for bullet in entry.get("bullets") or []:
                yield f"{section}/{entry.get('id')}/{bullet.get('id')}", bullet


def validate(data: dict[str, Any], *, require_sources: bool = False) -> list[str]:
    """Return a list of problems; empty means the base is well formed."""
    problems: list[str] = []
    seen: set[str] = set()

    def check_id(value: Any, where: str) -> None:
        if not isinstance(value, str) or not ID_RE.match(value):
            problems.append(f"{where}: id {value!r} must be lowercase letters, digits, '.' or '-'")
        elif value in seen:
            problems.append(f"{where}: duplicate id {value!r}")
        else:
            seen.add(value)

    check_id(data.get("id"), "base")
    for key in data.get("order") or []:
        if key not in SECTION_TITLES:
            problems.append(f"order: unknown section {key!r}")

    summary = data.get("summary")
    if summary is not None:
        if not isinstance(summary, dict) or not str(summary.get("text", "")).strip():
            problems.append("summary: needs a text field")
        else:
            check_id(summary.get("id"), "summary")

    for i, group in enumerate(data.get("skills") or []):
        where = f"skills[{i}]"
        check_id(group.get("id"), where)
        if not group.get("label") or not group.get("items"):
            problems.append(f"{where}: needs label and items")

    required = {
        "experience": ("title", "company", "dates"),
        "projects": ("name",),
        "education": ("school", "degree", "dates"),
    }
    for section, fields in required.items():
        for i, entry in enumerate(data.get(section) or []):
            where = f"{section}[{i}]"
            check_id(entry.get("id"), where)
            for field in fields:
                if not str(entry.get(field) or "").strip():
                    problems.append(f"{where}: missing {field}")

    for path, bullet in iter_bullets(data):
        check_id(bullet.get("id"), path)
        if not str(bullet.get("text") or "").strip():
            problems.append(f"{path}: empty text")
        sources = bullet.get("sources")
        if sources is not None and not isinstance(sources, list):
            problems.append(f"{path}: sources must be a list of passage ids")
        elif require_sources and not sources:
            problems.append(f"{path}: no source passage ids")
    return problems


def check_sources(data: dict[str, Any], index_path: Path) -> list[str]:
    """Check every cited passage id exists in the vault index and may back a resume claim."""
    import json

    passages = json.loads(Path(index_path).read_text(encoding="utf-8"))["passages"]
    use = {p["id"]: p.get("use", "claim") for p in passages}
    cited = [("summary", data.get("summary") or {})] + list(iter_bullets(data))
    problems = []
    for path, item in cited:
        for pid in item.get("sources") or []:
            if pid not in use:
                problems.append(f"{path}: unknown passage id {pid!r}")
            elif use[pid] != "claim":
                problems.append(f"{path}: passage {pid!r} is {use[pid]}, not a resume claim")
    return problems


# ---------------------------------------------------------------------------
# Low-level docx helpers
# ---------------------------------------------------------------------------


def _font(run, size: float, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def _paragraph(doc, *, space_before: float = 0, space_after: float = 0):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = 1.0
    return p


def _hyperlink(paragraph, url: str, text: str, size: float, *, bold: bool = False) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_pr.append(fonts)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    r_pr.append(color)
    for tag in ("w:sz", "w:szCs"):
        sz = OxmlElement(tag)
        sz.set(qn("w:val"), str(int(size * 2)))
        r_pr.append(sz)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _bottom_rule(paragraph) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)


def _url(value: str) -> str:
    return value if value.lower().startswith("http") else "https://" + value.lstrip("/")


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


class _Renderer:
    def __init__(self, data: dict[str, Any], contact: dict[str, str]):
        self.data = data
        self.contact = contact
        self.layout = Layout(**(data.get("layout") or {}))
        self.doc = Document()
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = section.right_margin = Inches(self.layout.side_margin_in)
        section.top_margin = Inches(self.layout.top_margin_in)
        section.bottom_margin = Inches(self.layout.bottom_margin_in)
        self.width = section.page_width - section.left_margin - section.right_margin
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT_NAME
        normal.font.size = Pt(self.layout.body_pt)

    @property
    def body(self) -> float:
        return self.layout.body_pt

    def build(self):
        self._header()
        for key in self.data.get("order") or DEFAULT_ORDER:
            if self.data.get(key):
                getattr(self, f"_{key}")()
        return self.doc

    def _header(self) -> None:
        p = _paragraph(self.doc)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(self.contact.get("name") or self.data.get("name", "")), self.layout.name_pt, bold=True)

        p = _paragraph(self.doc, space_after=2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        size = self.layout.contact_pt
        parts = [
            (k, self.contact.get(k, ""))
            for k in ("location", "email", "phone", "linkedin", "github")
        ]
        first = True
        for key, value in parts:
            if not value:
                continue
            if not first:
                _font(p.add_run(" | "), size)
            first = False
            if key == "email":
                _hyperlink(p, "mailto:" + value, value, size)
            elif key in ("linkedin", "github"):
                _hyperlink(p, _url(value), value, size)
            else:
                _font(p.add_run(value), size)

    def _heading(self, key: str) -> None:
        p = _paragraph(self.doc, space_before=self.layout.section_space_pt, space_after=2)
        _font(p.add_run(SECTION_TITLES[key].upper()), self.body + 0.5, bold=True)
        _bottom_rule(p)

    def _entry_line(self, runs: list[tuple[str, dict]], right: str, *, space_before: float, link: str = "") -> None:
        p = _paragraph(self.doc, space_before=space_before)
        p.paragraph_format.tab_stops.add_tab_stop(self.width, WD_TAB_ALIGNMENT.RIGHT)
        for i, (text, style) in enumerate(runs):
            if not text:
                continue
            if i == 0 and link:
                _hyperlink(p, _url(link), text, self.body, bold=style.get("bold", False))
            else:
                _font(p.add_run(text), self.body, **style)
        if right:
            p.add_run("\t")
            _font(p.add_run(right), self.body)

    def _bullets(self, bullets: list[dict]) -> None:
        for bullet in bullets:
            p = _paragraph(self.doc)
            fmt = p.paragraph_format
            fmt.left_indent = Inches(0.18)
            fmt.first_line_indent = Inches(-0.13)
            fmt.tab_stops.add_tab_stop(Inches(0.18))
            _font(p.add_run(f"{BULLET}\t{bullet['text'].strip()}"), self.body)

    def _summary(self) -> None:
        self._heading("summary")
        p = _paragraph(self.doc)
        _font(p.add_run(self.data["summary"]["text"].strip()), self.body)

    def _skills(self) -> None:
        self._heading("skills")
        for group in self.data["skills"]:
            p = _paragraph(self.doc)
            _font(p.add_run(f"{group['label']}: "), self.body, bold=True)
            _font(p.add_run(", ".join(str(x) for x in group["items"])), self.body)

    def _experience(self) -> None:
        self._heading("experience")
        for i, job in enumerate(self.data["experience"]):
            runs = [
                (job["title"], {"bold": True}),
                (", ", {}),
                (job["company"], {"bold": True}),
            ]
            if job.get("location"):
                runs.append((f", {job['location']}", {}))
            self._entry_line(runs, job["dates"], space_before=self.layout.entry_space_pt if i else 0)
            self._bullets(job.get("bullets") or [])

    def _projects(self) -> None:
        self._heading("projects")
        for i, proj in enumerate(self.data["projects"]):
            runs = [(proj["name"], {"bold": True})]
            if proj.get("context"):
                runs.append((f", {proj['context']}", {"italic": True}))
            self._entry_line(
                runs,
                str(proj.get("dates") or ""),
                space_before=self.layout.entry_space_pt if i else 0,
                link=proj.get("url", ""),
            )
            self._bullets(proj.get("bullets") or [])

    def _education(self) -> None:
        self._heading("education")
        for i, edu in enumerate(self.data["education"]):
            runs = [(edu["school"], {"bold": True})]
            if edu.get("location"):
                runs.append((f", {edu['location']}", {}))
            self._entry_line(runs, edu["dates"], space_before=self.layout.entry_space_pt if i else 0)
            p = _paragraph(self.doc)
            _font(p.add_run(edu["degree"].strip()), self.body)
            for line in edu.get("lines") or []:
                p = _paragraph(self.doc)
                label = line.get("label")
                if label:
                    _font(p.add_run(f"{label}: "), self.body, bold=True)
                _font(p.add_run(str(line["text"]).strip()), self.body)


def build_document(data: dict[str, Any], contact: dict[str, str]):
    return _Renderer(data, contact).build()


# ---------------------------------------------------------------------------
# PDF and checks
# ---------------------------------------------------------------------------


_SOFFICE_WINDOWS = (
    Path("C:/Program Files/LibreOffice/program/soffice.exe"),
    Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
)


def _word_to_pdf(docx_path: Path, pdf_path: Path) -> str | None:
    """Convert with Word through COM. Returns None on success, else why Word couldn't be used."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return "pywin32 isn't installed (run `uv sync`)"
    pythoncom.CoInitialize()
    try:
        try:
            word = win32com.client.DispatchEx("Word.Application")
        except Exception as e:  # noqa: BLE001 - pywin32 is installed but Word isn't
            return f"Word isn't available ({e})"
        try:
            word.Visible = False
            word.DisplayAlerts = 0
            doc = word.Documents.Open(str(docx_path), False, True, False)  # ConfirmConversions, ReadOnly, AddToRecentFiles
            try:
                # ExportAsFixedFormat is Word's PDF export; late-bound SaveAs isn't exposed by every
                # Word build ("AttributeError: Open.SaveAs"), so try the export calls in order.
                for call in (lambda: doc.ExportAsFixedFormat(str(pdf_path), 17),  # 17 = wdExportFormatPDF
                             lambda: doc.SaveAs2(str(pdf_path), FileFormat=17),
                             lambda: doc.SaveAs(str(pdf_path), FileFormat=17)):
                    try:
                        call()
                        return None
                    except Exception as e:  # noqa: BLE001 - try the next export call
                        err = e
                return f"Word couldn't export the PDF ({err})"
            finally:
                doc.Close(0)
        except Exception as e:  # noqa: BLE001 - any COM failure falls back to LibreOffice
            return f"Word failed ({e})"
        finally:
            word.Quit()
    finally:
        pythoncom.CoUninitialize()


def _find_soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    # The Windows installer doesn't put LibreOffice on PATH.
    return next((str(p) for p in _SOFFICE_WINDOWS if p.exists()), None) if sys.platform == "win32" else None


def docx_to_pdf(docx_path: Path, pdf_path: Path) -> Path:
    """Convert with Word (Windows) when it is installed, else LibreOffice."""
    docx_path, pdf_path = Path(docx_path).resolve(), Path(pdf_path).resolve()
    word_error = None
    if sys.platform == "win32":
        word_error = _word_to_pdf(docx_path, pdf_path)
        if word_error is None:
            return pdf_path
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError(
            "No PDF converter: " + (f"{word_error}, and " if word_error else "")
            + "LibreOffice isn't installed (Windows: `winget install TheDocumentFoundation.LibreOffice`)"
        )
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx_path)],
            check=True,
            capture_output=True,
        )
        shutil.move(str(Path(tmp) / (docx_path.stem + ".pdf")), pdf_path)
    return pdf_path


def pdf_page_count(pdf_path: Path) -> int:
    return len(re.findall(rb"/Type\s*/Page(?!s)", Path(pdf_path).read_bytes()))


def render(
    base_path: Path,
    out_dir: Path | None = None,
    *,
    stem: str | None = None,
    pdf: bool = True,
    contact: dict[str, str] | None = None,
) -> dict[str, Any]:
    """Render one base. Returns paths, page count and validation problems."""
    data = load_resume(base_path)
    problems = validate(data)
    if contact is None:
        contact = load_contact()
    # A base may override contact fields, e.g. ``contact: {location: "Vancouver, BC"}``.
    contact = {**contact, **{k: str(v) for k, v in (data.get("contact") or {}).items()}}
    out_dir = Path(out_dir or Path(base_path).parent)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = stem or data.get("file_stem") or Path(base_path).stem
    docx_path = out_dir / f"{stem}.docx"
    build_document(data, contact).save(str(docx_path))
    result: dict[str, Any] = {"docx": docx_path, "problems": problems}
    if pdf:
        pdf_path = docx_to_pdf(docx_path, out_dir / f"{stem}.pdf")
        result["pdf"] = pdf_path
        result["pages"] = pdf_page_count(pdf_path)
    return result
