"""Step 9: render the final resume and cover letter, fit the resume to one page, parse check both."""

from __future__ import annotations

import copy
import re
import tempfile
from datetime import date
from pathlib import Path
from typing import Any, Callable

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from src import resume as resume_mod
from src.tailor import TailorError
from src.tailor.appdir import _UNSAFE, load_app, log, read_json, write_json
from src.tailor.parsecheck import parse_check


def _safe(name: str) -> str:
    return re.sub(r"\s+", " ", _UNSAFE.sub(" ", name)).strip(" .")


# ---------------------------------------------------------------------------
# Sources
# ---------------------------------------------------------------------------


def _rounds(app_dir: Path) -> list[int]:
    found = []
    for p in (app_dir / "drafts").glob("r*"):
        m = re.fullmatch(r"r(\d+)", p.name)
        if m and p.is_dir():
            found.append(int(m.group(1)))
    return sorted(found)


def _source_path(app_dir: Path, app: dict[str, Any], name: str) -> Path:
    humanized = app_dir / "humanize" / name
    if humanized.exists():
        return humanized
    rounds = _rounds(app_dir)
    best = app.get("best_round")
    order = ([best] if best in rounds else []) + sorted(rounds, reverse=True)
    for rnd in order:
        path = app_dir / "drafts" / f"r{rnd}" / name
        if path.exists():
            return path
    raise TailorError(f"no {name} found in {app_dir} (looked in humanize/ and drafts/)")


# ---------------------------------------------------------------------------
# Page fit
# ---------------------------------------------------------------------------


def drop_order(data: dict[str, Any], relevance: dict[str, Any], protected: list[str]) -> list[str]:
    """Bullet paths in the order they may be dropped.

    Lowest relevance first; ties drop later bullets first. Protected bullets are never
    listed, and an entry always keeps at least one bullet. A bullet the edits added has no
    rating of its own (the plan rates base bullets and candidate passages), so it takes its
    best-rated cited passage's: unrated, it would always be the first to go.
    """
    positions: dict[str, int] = {}
    entry_of: dict[str, str] = {}
    score: dict[str, float] = {}
    for i, (path, b) in enumerate(resume_mod.iter_bullets(data)):
        positions[path] = i
        entry_of[path] = path.rsplit("/", 1)[0]
        if path in relevance:
            score[path] = float(relevance.get(path) or 0)
        else:
            score[path] = max((float(relevance.get(s) or 0) for s in b.get("sources") or []), default=0.0)
    remaining: dict[str, int] = {}
    for path, entry in entry_of.items():
        remaining[entry] = remaining.get(entry, 0) + 1
    candidates = [p for p in positions if p not in set(protected)]
    candidates.sort(key=lambda p: (score[p], -positions[p]))
    order = []
    for path in candidates:
        entry = entry_of[path]
        if remaining[entry] > 1:
            remaining[entry] -= 1
            order.append(path)
    return order


def without_bullets(data: dict[str, Any], paths: list[str]) -> dict[str, Any]:
    gone = set(paths)
    out = copy.deepcopy(data)
    for section in ("experience", "projects"):
        for entry in out.get(section) or []:
            entry["bullets"] = [
                b for b in entry.get("bullets") or [] if f"{section}/{entry.get('id')}/{b.get('id')}" not in gone
            ]
    return out


def fit_to_page(
    data: dict[str, Any],
    relevance: dict[str, Any],
    protected: list[str],
    pages_of: Callable[[dict[str, Any]], int],
) -> tuple[dict[str, Any], list[str], int]:
    """Drop the fewest bullets that make `pages_of(data)` <= 1 (binary search over the drop count).

    Returns (fitted data, dropped paths, pages of the result).
    """
    pages = pages_of(data)
    if pages <= 1:
        return data, [], pages
    order = drop_order(data, relevance, protected)
    lo, hi = 1, len(order)  # smallest k in [1, len(order)] that fits
    if not order:
        return data, [], pages
    best: tuple[int, int] | None = None
    last_pages = pages
    while lo <= hi:
        mid = (lo + hi) // 2
        p = pages_of(without_bullets(data, order[:mid]))
        if p <= 1:
            best, hi = (mid, p), mid - 1
        else:
            lo, last_pages = mid + 1, p
    if best is None:  # even dropping everything allowed does not fit
        return without_bullets(data, order), list(order), last_pages
    return without_bullets(data, order[: best[0]]), order[: best[0]], best[1]


# ---------------------------------------------------------------------------
# Cover letter
# ---------------------------------------------------------------------------


def build_letter_document(letter: dict[str, Any], contact: dict[str, Any], *, body_pt: float = 11,
                          today: str | None = None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = resume_mod.FONT_NAME
    normal.font.size = Pt(body_pt)

    p = resume_mod._paragraph(doc)
    resume_mod._font(p.add_run(contact.get("name", "")), 16, bold=True)
    p = resume_mod._paragraph(doc, space_after=14)
    first = True
    for key in ("location", "email", "phone", "linkedin", "github"):
        value = contact.get(key)
        if not value:
            continue
        if not first:
            resume_mod._font(p.add_run(" | "), 9.5)
        first = False
        if key == "email":
            resume_mod._hyperlink(p, "mailto:" + value, value, 9.5)
        elif key in ("linkedin", "github"):
            resume_mod._hyperlink(p, resume_mod._url(value), value, 9.5)
        else:
            resume_mod._font(p.add_run(value), 9.5)

    p = resume_mod._paragraph(doc, space_after=10)
    resume_mod._font(p.add_run(today or date.today().strftime("%B %d, %Y").replace(" 0", " ")), body_pt)
    if letter.get("greeting"):
        p = resume_mod._paragraph(doc, space_after=8)
        resume_mod._font(p.add_run(letter["greeting"].strip()), body_pt)
    for para in letter.get("paragraphs") or []:
        text = para.get("text", "") if isinstance(para, dict) else str(para)
        p = resume_mod._paragraph(doc, space_after=8)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        resume_mod._font(p.add_run(text.strip()), body_pt)
    lines = [ln.strip() for ln in str(letter.get("sign_off") or "").splitlines() if ln.strip()]
    for i, line in enumerate(lines):
        p = resume_mod._paragraph(doc, space_before=4 if i == 0 else (14 if i == 1 else 0))
        resume_mod._font(p.add_run(line), body_pt)
    return doc


def render_letter(letter: dict[str, Any], contact: dict[str, Any], out_dir: Path, stem: str,
                  *, today: str | None = None) -> dict[str, Any]:
    """Render the letter, shrinking the font a little if it spills onto a second page."""
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path, pdf_path = out_dir / f"{stem}.docx", out_dir / f"{stem}.pdf"
    pages = 0
    for pt in (11, 10.5, 10):
        build_letter_document(letter, contact, body_pt=pt, today=today).save(str(docx_path))
        resume_mod.docx_to_pdf(docx_path, pdf_path)
        pages = resume_mod.pdf_page_count(pdf_path)
        if pages <= 1:
            break
    return {"docx": docx_path, "pdf": pdf_path, "pages": pages}


# ---------------------------------------------------------------------------
# Step 9
# ---------------------------------------------------------------------------


def _contact(cfg: dict[str, Any], data: dict[str, Any]) -> dict[str, str]:
    if cfg.get("contact"):
        base = {k: str(v) for k, v in cfg["contact"].items()}
    else:
        try:
            base = resume_mod.load_contact()
        except FileNotFoundError:
            base = {}
    base.update({k: str(v) for k, v in (data.get("contact") or {}).items()})
    if not base.get("name"):
        raise TailorError("no contact name: set contact.name in the resume yaml or config/contact.yaml")
    return base


def _company_twins(app_dir: Path, app: dict[str, Any]) -> bool:
    """True when another application folder beside this one is for the same company."""
    me = str(app.get("company") or "").strip().lower()
    for other in app_dir.parent.glob("*/app.json"):
        if other.parent.resolve() == app_dir.resolve():
            continue
        try:
            if str(read_json(other).get("company") or "").strip().lower() == me:
                return True
        except Exception:  # noqa: BLE001 - a broken sibling folder doesn't stop this render
            continue
    return False


def run_render(app_dir: Path, cfg: dict[str, Any]) -> dict[str, Any]:
    app_dir = Path(app_dir)
    app = load_app(app_dir)
    company = _safe(str(app.get("company") or "")) or "Company"
    resume_src = _source_path(app_dir, app, "resume.yaml")
    letter_src = _source_path(app_dir, app, "cover-letter.json")
    data = yaml.safe_load(resume_src.read_text(encoding="utf-8")) or {}
    letter = read_json(letter_src)
    contact = _contact(cfg, data)
    name = _safe(contact["name"])

    plan_path = app_dir / "plan.json"
    plan = read_json(plan_path) if plan_path.exists() else {}
    relevance = plan.get("bullet_relevance") or {}
    protected = plan.get("protected_paths") or []

    final = app_dir / "final"
    final.mkdir(exist_ok=True)
    # Two applications to one company would upload as identical file names, so add the job title.
    if _company_twins(app_dir, app):
        company = f"{company} - {_safe(str(app.get('title') or ''))[:50].strip()}".rstrip(" -")
    resume_stem = f"{name} Resume - {company}"
    letter_stem = f"{name} Cover Letter - {company}"

    def pages_of(candidate: dict[str, Any]) -> int:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "try.yaml"
            path.write_text(yaml.safe_dump(candidate, sort_keys=False, allow_unicode=True), encoding="utf-8")
            return resume_mod.render(path, Path(tmp), stem="try", contact=contact)["pages"]

    fitted, dropped, _ = fit_to_page(data, relevance, protected, pages_of)

    yaml_path = final / "resume.yaml"
    yaml_path.write_text(yaml.safe_dump(fitted, sort_keys=False, allow_unicode=True), encoding="utf-8")
    res = resume_mod.render(yaml_path, final, stem=resume_stem, contact=contact)
    write_json(final / "cover-letter.json", letter)
    let = render_letter(letter, contact, final, letter_stem)

    resume_check = parse_check(res["pdf"], {**fitted, "kind": "resume"}, contact)
    letter_check = parse_check(let["pdf"], {**letter, "kind": "letter"}, contact)
    ok = resume_check["ok"] and letter_check["ok"]
    result = {
        "ok": ok,
        "resume": {**resume_check, "pdf": res["pdf"].name, "docx": res["docx"].name, "pages": res["pages"]},
        "letter": {**letter_check, "pdf": let["pdf"].name, "docx": let["docx"].name, "pages": let["pages"]},
        "dropped": dropped,
        "sources": {"resume": str(resume_src.relative_to(app_dir)), "letter": str(letter_src.relative_to(app_dir))},
    }
    write_json(final / "parse.json", result)
    note = f"rendered; parse {'ok' if ok else 'FAILED'}; dropped {len(dropped)} bullet(s)"
    log(app_dir, "render", note, **({} if ok else {"status": "needs_fixes"}))
    return result
