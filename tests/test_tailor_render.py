"""Tests for company.py, parsecheck.py and render.py (steps 1b and 9)."""

from __future__ import annotations

import json
import os
import shutil
import time
from pathlib import Path

import pytest
import yaml

from src import resume as resume_mod
from src.tailor import TailorError
from src.tailor.company import company_dir, company_slug, is_stale, verify_company
from src.tailor.parsecheck import parse_check
from src.tailor.render import build_letter_document, drop_order, fit_to_page, run_render

FIX = Path(__file__).parent / "fixtures" / "tailor"
needs_soffice = pytest.mark.skipif(
    not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed"
)
CONTACT = {"name": "Alex Rivera", "location": "Vancouver, BC", "email": "alex@example.com",
           "phone": "604-555-0100", "linkedin": "", "github": ""}


def _base() -> dict:
    return yaml.safe_load((FIX / "base-analyst.yaml").read_text(encoding="utf-8"))


# ----- company -----


def test_slug_and_dir(tmp_path):
    assert company_slug("  Acme, Inc. & Sons! ") == "acme-inc-sons"
    # tmp_path is absolute on every OS; "/tmp" gains a drive letter on Windows.
    cfg = {"paths": {"companies": str(tmp_path / "cos")}}
    assert company_dir(cfg, "Acme Co") == tmp_path / "cos" / "acme-co"


def test_verify_company(tmp_path):
    d = tmp_path / "co"
    shutil.copytree(FIX / "render", d)
    (d / "notes.md").write_text("research notes", encoding="utf-8")
    out = verify_company(d)
    by_id = {f["id"]: f for f in out["facts"]}
    assert by_id["f1"]["verified"] and by_id["f2"]["verified"]  # hyphenation, curly quotes
    assert not by_id["f3"]["verified"] and "not found" in by_id["f3"]["reason"]
    assert not by_id["f4"]["verified"] and "9.txt" in by_id["f4"]["reason"]
    assert (out["verified"], out["unverified"]) == (2, 2)
    saved = json.loads((d / "facts.json").read_text(encoding="utf-8"))
    assert len(saved) == 4 and saved[2]["verified"] is False
    assert (d / "notes.md").read_text(encoding="utf-8") == "research notes"  # research notes kept
    notes = (d / "verified.md").read_text(encoding="utf-8")
    assert "Founded 2015" in notes and "https://example.com/about" in notes
    assert notes.index("Unverified") < notes.index("Has 400 staff")


def test_verify_company_malformed(tmp_path):
    facts = tmp_path / "facts.json"
    for bad in ('[{"id": "f1", "fact": "x"}]', '{"a": 1}', "nope"):
        facts.write_text(bad, encoding="utf-8")
        with pytest.raises(TailorError):
            verify_company(tmp_path)


def test_is_stale(tmp_path):
    assert is_stale(tmp_path, 30)
    f = tmp_path / "facts.json"
    f.write_text("[]")
    assert not is_stale(tmp_path, 30)
    old = time.time() - 40 * 86400
    os.utime(f, (old, old))
    assert is_stale(tmp_path, 30)


# ----- page fit -----


def _long() -> dict:
    data = _base()
    data["experience"][0]["bullets"] = [
        {"id": f"northwind.b{i}", "text": f"Bullet number {i}"} for i in range(5)
    ]
    return data


def test_drop_order_rules():
    data = _long()
    n = "experience/northwind/northwind.b"
    rel = {f"{n}{i}": 1 for i in range(5)}
    rel[n + "1"] = 0
    order = drop_order(data, rel, [n + "4"])
    assert order == [n + "1", n + "3", n + "2", n + "0"]  # lowest first, then later first
    # lab and projects have one bullet each: never dropped


def test_drop_order_rates_an_added_bullet_by_its_sources():
    # A bullet a repair round appended isn't in the plan's ratings; it takes its cited passage's
    # rating instead of 0, so it isn't cut ahead of weaker base bullets.
    data = _long()
    n = "experience/northwind/northwind."
    data["experience"][0]["bullets"].append({"id": "northwind.new-1", "text": "Added", "sources": ["nw.talk"]})
    rel = {f"{n}b{i}": 1 for i in range(5)}
    rel["nw.talk"] = 2
    assert drop_order(data, rel, []) == [n + f"b{i}" for i in (4, 3, 2, 1, 0)]


def test_fit_to_page_binary_search():
    data = _long()

    def pages_of(d):
        n = sum(len(e["bullets"]) for e in d["experience"])
        return 1 if n <= 5 else 2  # lab 1 + northwind 4 fits

    fitted, dropped, pages = fit_to_page(data, {}, [], pages_of)
    assert pages == 1
    assert dropped == ["experience/northwind/northwind.b4"]  # tie: later bullet first
    assert len(fitted["experience"][0]["bullets"]) == 4
    assert len(data["experience"][0]["bullets"]) == 5  # input untouched


def test_fit_to_page_cannot_fit():
    data = _long()
    keep = "experience/northwind/northwind.b0"
    fitted, dropped, pages = fit_to_page(data, {}, [keep], lambda d: 2)
    assert pages == 2 and keep not in dropped
    assert [b["id"] for b in fitted["experience"][0]["bullets"]] == ["northwind.b0"]


def test_fit_to_page_already_fits():
    fitted, dropped, pages = fit_to_page(_long(), {}, [], lambda d: 1)
    assert dropped == [] and pages == 1


# ----- render + parse check (need LibreOffice) -----


@needs_soffice
def test_parse_check_on_rendered_resume(tmp_path):
    res = resume_mod.render(FIX / "base-analyst.yaml", tmp_path, contact=CONTACT)
    source = {**_base(), "kind": "resume"}
    out = parse_check(res["pdf"], source, CONTACT)
    assert out["ok"], out
    assert {c["code"] for c in out["checks"]} >= {"text_extracts", "one_page", "headings", "round_trip"}

    source["experience"][0]["bullets"][0]["text"] = "Negotiated vendor contracts across three continents."
    bad = parse_check(res["pdf"], source, CONTACT)
    assert not bad["ok"]
    assert bad["round_trip_failures"][0]["reason"] == "missing"

    wrong = parse_check(res["pdf"], {**_base(), "kind": "resume"}, {**CONTACT, "email": "other@nowhere.org"})
    assert not wrong["ok"]


@needs_soffice
def test_run_render_end_to_end(tmp_path):
    app = tmp_path / "app"
    r1 = app / "drafts" / "r1"
    r1.mkdir(parents=True)
    (r1 / "resume.yaml").write_text((FIX / "base-analyst.yaml").read_text(encoding="utf-8"), encoding="utf-8")
    letter = {
        "greeting": "Dear Hiring Manager,",
        "paragraphs": [
            {"text": "I wrote SQL daily at a logistics software company and enjoyed turning requests into dashboards."},
            {"text": "I would like to bring that habit of tracing every metric back to its source to your team."},
        ],
        "sign_off": "Sincerely,\nAlex Rivera",
    }
    (r1 / "cover-letter.json").write_text(json.dumps(letter), encoding="utf-8")
    (app / "app.json").write_text(
        json.dumps({"company": "Acme Corp", "status": "drafted", "best_round": 1}), encoding="utf-8"
    )
    # Files from an earlier render under another name don't survive a rerender.
    (app / "final").mkdir()
    (app / "final" / "Alex Rivera Resume - Acme Corp - Old Ti.pdf").write_bytes(b"old")
    out = run_render(app, {"contact": CONTACT})
    assert not (app / "final" / "Alex Rivera Resume - Acme Corp - Old Ti.pdf").exists()
    assert out["ok"], out
    final = app / "final"
    for name in ("Alex Rivera Resume - Acme Corp.pdf", "Alex Rivera Cover Letter - Acme Corp.docx",
                 "resume.yaml", "cover-letter.json", "parse.json"):
        assert (final / name).exists(), name
    app_json = json.loads((app / "app.json").read_text(encoding="utf-8"))
    assert app_json["status"] == "drafted" and app_json["history"][-1]["step"] == "render"


def test_page_text_drops_chrome():
    from src.tailor.company import page_text

    html = ("<html><body><nav>Menu Careers</nav><main><h1>About</h1><p>We deliver   care to "
            "1.25 million people.</p><script>var x=1;</script></main><footer>Copyright</footer></body></html>")
    assert page_text(html) == "About\nWe deliver care to 1.25 million people."


def test_page_text_keeps_inline_elements_in_their_sentence(tmp_path):
    from src.tailor.company import page_text

    html = ("<main><p>Acme was named one of\n<a href='/x'>Ontario's Top Garden Suppliers for 2025</a>, its third "
            "year.</p><p>We have placed <span>1</span><span>20+ gardeners</span> since 2019.<br>Call us.</p></main>")
    assert page_text(html) == ("Acme was named one of Ontario's Top Garden Suppliers for 2025, its third year.\n"
                               "We have placed 120+ gardeners since 2019.\nCall us.")
    d = tmp_path / "acme"
    (d / "pages").mkdir(parents=True)
    (d / "pages" / "1.txt").write_text("URL: x\n\n" + page_text(html), encoding="utf-8")
    (d / "facts.json").write_text(json.dumps([
        {"id": "f1", "fact": "award", "quote": "one of Ontario's Top Garden Suppliers for 2025, its third year",
         "page": 1},
        {"id": "f2", "fact": "placements", "quote": "placed 120+ gardeners", "page": 1}]), encoding="utf-8")
    assert verify_company(d)["verified"] == 2


def test_company_research_prompt_names_its_search_tool():
    # A subagent told only "search the web" reached for a search skill instead and stalled.
    root = Path(__file__).parent.parent
    prompt = (root / "prompts" / "tailor" / "company-research.md").read_text(encoding="utf-8")
    assert "WebSearch tool" in prompt and "Do not use skills" in prompt
    assert "WebSearch tool (no skills)" in (root / "docs" / "tailoring-runbook.md").read_text(encoding="utf-8")


def test_runbook_says_to_shorten_a_long_bullet_and_reapply():
    # apply-edits refuses a bullet over the limit so the run shortens it in the same round; a
    # runbook saying every rejection is "normal, continue" lost the bullet for that round instead.
    from src.tailor.edits import BULLET_MAX

    root = Path(__file__).parent.parent
    runbook = (root / "docs" / "tailoring-runbook.md").read_text(encoding="utf-8")
    assert '"shorten it"' in runbook and f"over {BULLET_MAX} characters" in runbook
    edits = (root / "prompts" / "tailor" / "resume-edits.md").read_text(encoding="utf-8")
    assert "sent back for repair" not in edits and "rerun\n   `apply-edits` for the same round" in edits


def test_fetch_page_saves_next_page(tmp_path, monkeypatch):
    import httpx

    from src.tailor.company import fetch_page

    def fake_get(url, **kw):
        return httpx.Response(200, headers={"content-type": "text/html"},
                              text="<main><p>Healthy lives in healthy communities.</p></main>",
                              request=httpx.Request("GET", url))

    monkeypatch.setattr(httpx, "get", fake_get)
    (tmp_path / "pages").mkdir()
    (tmp_path / "pages" / "1.txt").write_text("URL: x", encoding="utf-8")
    out = fetch_page(tmp_path, "https://example.com/about")
    assert out["page"] == 2
    saved = (tmp_path / "pages" / "2.txt").read_text(encoding="utf-8")
    assert saved.startswith("URL: https://example.com/about\n") and "Healthy lives in healthy communities." in saved


def test_fetch_page_rejects_a_page_of_loading_placeholders(tmp_path, monkeypatch):
    import httpx

    from src.tailor import TailorError
    from src.tailor.company import fetch_page

    html = {"/careers": "<main><h1>Careers</h1>" + "<div>Loading component...</div>" * 4 + "</main>",
            "/about": "<main><h1>About us</h1><p>We sell garden tools in Ottawa.</p><div>Loading...</div></main>"}
    monkeypatch.setattr(httpx, "get", lambda url, **kw: httpx.Response(
        200, headers={"content-type": "text/html"}, text=html[url[len("https://example.com"):]],
        request=httpx.Request("GET", url)))
    with pytest.raises(TailorError, match="no readable text"):
        fetch_page(tmp_path, "https://example.com/careers")
    assert not (tmp_path / "pages").exists()
    assert fetch_page(tmp_path, "https://example.com/about")["page"] == 1  # real text with one placeholder is kept


def test_fetch_page_rejects_pdf(tmp_path, monkeypatch):
    import httpx

    from src.tailor import TailorError
    from src.tailor.company import fetch_page

    monkeypatch.setattr(httpx, "get", lambda url, **kw: httpx.Response(
        200, headers={"content-type": "application/pdf"}, content=b"%PDF", request=httpx.Request("GET", url)))
    with pytest.raises(TailorError, match="not a web page"):
        fetch_page(tmp_path, "https://example.com/a.pdf")


def test_phrase_patterns_flag_dressed_up_lines():
    from src.tailor import phrases

    def codes(t, posting=""):
        return {h.code for h in phrases.scan(t, posting)}

    assert "deep_saying" in codes("Health care is about as real as problems get.")
    assert "fancy_verb" in codes("The portal serves as a hub for patients.")
    assert "run_up" in codes("Here's the thing. I like data.")
    assert codes("I wrote SQL for the marketing team.") == set()
    # protection rule: an exact posting phrase is allowed
    assert "inflated" not in codes("Data plays a key role here.", posting="Data plays a key role at VCH")


def test_triple_needs_a_real_list_of_three():
    from src.tailor import phrases

    def triples(t):
        return [h.phrase for h in phrases.scan(t, letter=True) if h.code == "triple"]

    assert triples("Our team is fast, careful and kind.") == ["is fast, careful and kind"]
    # An opening clause, then "I" and two verbs, is one comma and one "and" but not a list.
    assert triples("After hours, I built and tested a scheduling app.") == []
    assert triples("At the clinic, we planned and ran the survey.") == []
    # Naming real tools is content, not rhythm.
    assert triples("I checked data from Shopify, Yotpo and Braze.") == []
    assert triples("Data from Contoso, Fabrikam and others landed in the warehouse.") == []
    assert triples("Data from Contoso, Fabrikam and other vendors landed there.") == []
    # A clause, then a phrase starting with a preposition, is not a list of three.
    assert triples("It loads orders into SQLite, with retries and 40 tests.") == []
    assert triples("I want to keep the system running, by testing and fixing its jobs.") == []
    assert triples("We are quick, careful and others notice.") == ["are quick, careful and others"]
    # A list quoted from the posting is protected, with or without its Oxford comma.
    posting = "Our work on seat belts, air bags, and graduated driving laws saved lives."
    hits = [h for h in phrases.scan("I admire your work on seat belts, air bags and graduated driving laws.",
                                    posting, letter=True) if h.code == "triple"]
    assert hits == []


def test_fetch_page_refuses_job_boards_and_profile_sites(tmp_path, monkeypatch):
    import httpx

    from src.tailor import TailorError
    from src.tailor.company import fetch_page

    def fail(url, **kw):
        raise AssertionError("should not download")

    monkeypatch.setattr(httpx, "get", fail)
    for url in ("https://ca.indeed.com/cmp/acme", "https://www.crunchbase.com/organization/acme",
                "https://www.linkedin.com/company/acme"):
        with pytest.raises(TailorError, match="job board, profile or social site"):
            fetch_page(tmp_path, url)


def test_review_directory_facts_are_not_verified(tmp_path, monkeypatch):
    # A clinic directory's star rating verified as a company fact (Oct 25 2026): fetch refuses
    # review directories, and verify drops a fact from such a page even when it was saved earlier.
    import httpx

    from src.tailor.company import fetch_page

    monkeypatch.setattr(httpx, "get", lambda url, **kw: (_ for _ in ()).throw(AssertionError("no download")))
    with pytest.raises(TailorError, match="job board, profile or social site"):
        fetch_page(tmp_path, "https://medimap.ca/clinic/acme-clinic")
    pages = tmp_path / "pages"
    pages.mkdir()
    (pages / "1.txt").write_text("URL: https://www.ratemds.com/clinic/acme\n\n4.9/5 Average Rating\n", encoding="utf-8")
    (pages / "2.txt").write_text("URL: https://acme.example/about\n\nAcme runs three clinics.\n", encoding="utf-8")
    (tmp_path / "facts.json").write_text(json.dumps([
        {"id": "f1", "fact": "Rated 4.9", "quote": "4.9/5 Average Rating", "page": 1, "read": "2026-10-25"},
        {"id": "f2", "fact": "Three clinics", "quote": "Acme runs three clinics", "page": 2, "read": "2026-10-25"},
        {"id": "f3", "fact": "Three clinics", "quote": "Acme runs three clinics", "page": 2, "read": "2026-10-25",
         "url": "https://www.yelp.com/biz/acme"},
    ]), encoding="utf-8")
    by_id = {f["id"]: f for f in verify_company(tmp_path)["facts"]}
    assert not by_id["f1"]["verified"] and "ratemds" in by_id["f1"]["reason"]
    assert by_id["f2"]["verified"]
    assert not by_id["f3"]["verified"] and "yelp" in by_id["f3"]["reason"]


def test_fetch_page_stops_at_the_page_budget(tmp_path, monkeypatch):
    import httpx

    from src.tailor import TailorError
    from src.tailor.company import PAGE_BUDGET, fetch_page

    monkeypatch.setattr(httpx, "get", lambda url, **kw: httpx.Response(
        200, headers={"content-type": "text/html"}, text="<main><p>We sell garden tools.</p></main>",
        request=httpx.Request("GET", url)))
    for _ in range(PAGE_BUDGET):
        fetch_page(tmp_path, "https://example.com/about")
    with pytest.raises(TailorError, match="page budget used"):
        fetch_page(tmp_path, "https://example.com/news")


def test_old_news_is_not_a_verified_fact(tmp_path, monkeypatch):
    import httpx

    from src.tailor.company import fetch_page

    html = ('<head><meta property="article:published_time" content="2020-05-01T09:00:00Z"></head>'
            "<main><p>Acme acquired Widgetco to grow its analytics practice.</p></main>")
    monkeypatch.setattr(httpx, "get", lambda url, **kw: httpx.Response(
        200, headers={"content-type": "text/html"}, text=html, request=httpx.Request("GET", url)))
    out = fetch_page(tmp_path, "https://news.example.com/acme")
    assert out["published"] == "2020-05-01"
    (tmp_path / "pages" / "2.txt").write_text("URL: https://acme.example/about\n\nAcme sells tools.\n",
                                              encoding="utf-8")
    (tmp_path / "facts.json").write_text(json.dumps([
        {"id": "f1", "fact": "Bought Widgetco", "quote": "Acme acquired Widgetco", "page": 1, "read": "2026-10-21"},
        {"id": "f2", "fact": "Sells tools", "quote": "Acme sells tools", "page": 2, "read": "2026-10-21"},
    ]), encoding="utf-8")
    by_id = {f["id"]: f for f in verify_company(tmp_path)["facts"]}
    assert not by_id["f1"]["verified"] and "old news" in by_id["f1"]["reason"]
    assert by_id["f2"]["verified"]  # an undated page (the company's own) is fine


def test_company_twins_finds_a_second_application_to_the_same_company(tmp_path):
    from src.tailor.render import _company_twins

    def app(folder, company):
        d = tmp_path / folder
        d.mkdir()
        (d / "app.json").write_text(json.dumps({"company": company}), encoding="utf-8")
        return d

    ds = app("2026-10-20 Amazon - Data Scientist", "Amazon")
    app("2026-10-20 Amazon - BI Engineer", "amazon")
    solo = app("2026-10-20 Allstate - Analyst", "Allstate")
    assert _company_twins(ds, {"company": "Amazon"})
    assert not _company_twins(solo, {"company": "Allstate"})


def test_file_name_title_keeps_whole_words_and_drops_brackets():
    from src.tailor.render import _short_title

    assert _short_title("Business Intelligence Engineer, IT PnD Data & Analytics (DnA)") == \
        "Business Intelligence Engineer, IT PnD Data & Analytics"
    assert _short_title("Data Engineer (12 Month Contract)") == "Data Engineer"
    assert _short_title("Senior Data-Driven Experimentation and Machine Learning Platform Engineer for Payments") == \
        "Senior Data-Driven Experimentation and Machine Learning"


def test_letter_heading_matches_resume():
    """The letter reuses the resume's centered heading, as wide as the resume's text block."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    doc = build_letter_document({"greeting": "Hi,", "paragraphs": ["Body."]}, CONTACT, today="October 1, 2026")
    name, line = doc.paragraphs[0], doc.paragraphs[1]
    assert name.text == "Alex Rivera"
    assert "604-555-0100" in line.text
    layout = resume_mod.Layout()
    for p in (name, line):
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert p.paragraph_format.left_indent == -Inches(1 - layout.side_margin_in)
    assert name.runs[0].font.size.pt == layout.name_pt
